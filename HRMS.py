import streamlit as st
import pandas as pd
from datetime import datetime, date, timedelta
from bson import ObjectId
import smtplib
import ssl
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pymongo import MongoClient, ASCENDING

# ====================== CONFIG ======================
st.set_page_config(page_title="Kaunda HRMS", layout="wide", page_icon="🧑‍💼")
st.title("🧑‍💼 Kaunda HR Management System")

# ====================== MONGODB CONNECTION ======================
# The connection string is stored securely in .streamlit/secrets.toml
# It is NEVER written in this code file.
@st.cache_resource
def get_db():
    uri    = st.secrets["mongodb"]["uri"]
    client = MongoClient(uri)
    return client["HRMS"]

db = get_db()

# ---- Collections (equivalent to SQLite tables) ----
employees_col        = db["employees"]
payroll_settings_col = db["payroll_settings"]
payroll_runs_col     = db["payroll_runs"]
payroll_payments_col = db["payroll_payments"]
leave_col            = db["leave_tracker"]
disciplinary_col     = db["disciplinary_tracker"]
settings_col         = db["system_settings"]
notif_col            = db["notification_log"]
alert_rules_col      = db["alert_rules"]

def init_db():
    if alert_rules_col.count_documents({}) == 0:
        alert_rules_col.insert_many([
            {"rule_name":"Contract Expiry Warning",        "rule_type":"contract_expiry",   "threshold_days":30, "enabled":True},
            {"rule_name":"Probation Period End",           "rule_type":"probation_end",      "threshold_days":14, "enabled":True},
            {"rule_name":"Payroll Processed Notification", "rule_type":"payroll_processed",  "threshold_days":0,  "enabled":True},
            {"rule_name":"New Employee Added",             "rule_type":"new_employee",        "threshold_days":0,  "enabled":True},
        ])

init_db()

# ====================== SETTINGS HELPERS ======================
def get_setting(key, default=""):
    doc = settings_col.find_one({"key": key})
    return doc["value"] if doc else default

def save_setting(key, value):
    settings_col.update_one({"key": key}, {"$set": {"value": value}}, upsert=True)

# ====================== LOGIN ======================
def init_login_session():
    if "logged_in" not in st.session_state: st.session_state.logged_in = False
    if "user_role" not in st.session_state: st.session_state.user_role = None
    if "username"  not in st.session_state: st.session_state.username  = None

def show_login_page():
    st.markdown("""<style>
    .stButton > button {
        background:linear-gradient(135deg,#4CAF50,#388E3C);
        color:white;border:none;border-radius:8px;
        padding:12px 24px;font-size:16px;font-weight:bold;width:100%;
    }</style>""", unsafe_allow_html=True)
    col1,col2,col3 = st.columns([1,2,1])
    with col2:
        st.markdown('<h2 style="color:#2E7D32;text-align:center">🔐 Admin Login</h2>', unsafe_allow_html=True)
        admin_user = get_setting("admin_username","admin")
        admin_pass = get_setting("admin_password","admin123")
        with st.form("login_form"):
            username = st.text_input("Username")
            password = st.text_input("Password", type="password")
            if st.form_submit_button("Login", use_container_width=True):
                if username == admin_user and password == admin_pass:
                    st.session_state.logged_in = True
                    st.session_state.user_role = "admin"
                    st.session_state.username  = username
                    st.success("✅ Login successful!")
                    st.rerun()
                else:
                    st.error("❌ Invalid credentials.")

# ====================== EMAIL ======================
def send_email(subject, body, recipient=None, html_body=None):
    smtp_host   = get_setting("smtp_host","")
    smtp_port   = get_setting("smtp_port","587")
    smtp_user   = get_setting("smtp_user","")
    smtp_pass   = get_setting("smtp_password","")
    from_name   = get_setting("company_name","Kaunda HRMS")
    admin_email = get_setting("admin_email","")
    to_email    = recipient or admin_email
    if not all([smtp_host, smtp_user, smtp_pass, to_email]):
        return False, "Email settings not fully configured."
    try:
        msg = MIMEMultipart("alternative")
        msg["Subject"] = subject
        msg["From"]    = f"{from_name} <{smtp_user}>"
        msg["To"]      = to_email
        msg.attach(MIMEText(body,"plain"))
        if html_body: msg.attach(MIMEText(html_body,"html"))
        context = ssl.create_default_context()
        port    = int(smtp_port)
        if port == 465:
            with smtplib.SMTP_SSL(smtp_host, port, context=context) as s:
                s.login(smtp_user, smtp_pass)
                s.sendmail(smtp_user, to_email, msg.as_string())
        else:
            with smtplib.SMTP(smtp_host, port) as s:
                s.ehlo(); s.starttls(context=context)
                s.login(smtp_user, smtp_pass)
                s.sendmail(smtp_user, to_email, msg.as_string())
        notif_col.insert_one({"sent_at":datetime.now(),"subject":subject,"recipient":to_email,"body":body,"status":"sent"})
        return True, "Email sent successfully!"
    except Exception as e:
        notif_col.insert_one({"sent_at":datetime.now(),"subject":subject,"recipient":to_email,"body":body,"status":f"failed:{e}"})
        return False, f"Email error: {e}"

def build_html_email(title, content_rows, footer=""):
    company   = get_setting("company_name","Kaunda HRMS")
    rows_html = "".join(
        f"<tr><td style='padding:6px 12px;border-bottom:1px solid #eee;font-weight:bold;width:40%;color:#555'>{k}</td>"
        f"<td style='padding:6px 12px;border-bottom:1px solid #eee;color:#222'>{v}</td></tr>"
        for k,v in content_rows)
    return f"""<html><body style='font-family:Arial,sans-serif;background:#f4f6f8;margin:0;padding:20px'>
    <div style='max-width:600px;margin:auto;background:#fff;border-radius:8px;overflow:hidden;box-shadow:0 2px 8px rgba(0,0,0,.1)'>
      <div style='background:#2E7D32;padding:20px 30px;color:white'>
        <h2 style='margin:0'>🧑‍💼 {company}</h2><p style='margin:4px 0 0;font-size:14px;opacity:.9'>{title}</p>
      </div>
      <div style='padding:24px 30px'><table style='width:100%;border-collapse:collapse'>{rows_html}</table>
        {f'<p style="margin-top:20px;color:#777;font-size:13px">{footer}</p>' if footer else ''}
      </div>
      <div style='background:#f0f0f0;padding:12px 30px;font-size:12px;color:#999;text-align:center'>
        Sent by {company} HRMS &bull; {datetime.now().strftime('%d %b %Y %H:%M')}
      </div></div></body></html>"""

# ====================== EXCEL ======================
def export_to_excel(dataframes_dict, filename):
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for sheet_name, df in dataframes_dict.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            for cell in ws[1]:
                cell.font      = Font(bold=True, color="FFFFFF")
                cell.fill      = PatternFill(start_color="2E7D32",end_color="2E7D32",fill_type="solid")
                cell.alignment = Alignment(horizontal="center",vertical="center")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(max((len(str(c.value or "")) for c in col),default=0)+2, 50)
    output.seek(0)
    return output

# ====================== MONGO HELPERS ======================
def docs_to_df(docs):
    if not docs: return pd.DataFrame()
    df = pd.DataFrame(docs)
    if "_id" in df.columns: df["_id"] = df["_id"].astype(str)
    return df

def get_employees():
    return docs_to_df(list(employees_col.find().sort("name", ASCENDING)))

def get_today(): return datetime.now().date()

def get_current_period():
    now = datetime.now(); return now.year, now.month

def calculate_pay(basic, housing, transport, other_all, tax_rate, pension_rate, other_ded):
    allowances = housing + transport + other_all
    gross = basic + allowances
    tax   = gross * tax_rate
    pension = gross * pension_rate
    net = gross - (tax + pension + other_ded)
    return {"basic_salary":basic,"allowances":allowances,"gross_pay":gross,
            "tax":tax,"pension":pension,"other_deduction":other_ded,"net_pay":net}

def generate_employee_id(department):
    dept_code = department[:3].upper() if department else "EMP"
    count     = employees_col.count_documents({}) + 1
    return f"EMP-{dept_code}-{count:03d}"

def contract_status(contract_end, actual_end_date, threshold=30):
    if actual_end_date: return "LEFT"
    if not contract_end: return "UNKNOWN"
    try:
        end  = datetime.strptime(str(contract_end)[:10],"%Y-%m-%d").date()
        days = (end - get_today()).days
        if days < 0: return "EXPIRED"
        elif days <= threshold: return "EXPIRING SOON"
        else: return "ACTIVE"
    except: return "UNKNOWN"

def disciplinary_status(expiry_date):
    if not expiry_date: return "NO EXPIRY"
    try:
        exp  = datetime.strptime(str(expiry_date)[:10],"%Y-%m-%d").date()
        days = (exp - get_today()).days
        if days < 0: return "EXPIRED"
        elif days <= 7: return "EXPIRING SOON"
        else: return "ACTIVE"
    except: return "UNKNOWN"

DEPARTMENTS     = ["Human Resources","Finance","Information Technology","Operations","Sales & Marketing",
                   "Administration","Legal","Procurement","Customer Service","Engineering","Mechanical",
                   "Electrical","General Services","Stores","Logistics","Management","Other"]
CONTRACT_TYPES  = ["Permanent","Fixed-Term","Contract","Part-Time","Internship","Probation"]
STATUS_OPTIONS  = ["active","inactive","suspended","terminated","on_leave"]
LEAVE_TYPES     = ["Annual Leave","Sick Leave","Maternity Leave","Paternity Leave",
                   "Compassionate Leave","Study Leave","Unpaid Leave","Other"]
APPROVAL_STATUS = ["Pending","Approved","Rejected","Cancelled"]
ISSUE_TYPES     = ["Verbal Warning","Written Warning","First Warning","Final Warning",
                   "Suspension","Misconduct","Insubordination","Absenteeism","Other"]
SEPARATION_TYPES = ["Resigned","Terminated","Contract Ended","Retired","Deceased","Other"]

# ====================== BOOT ======================
init_login_session()
if not st.session_state.logged_in:
    show_login_page()
    st.stop()

# ====================== SIDEBAR ======================
st.sidebar.markdown(f"**Logged in as:** {st.session_state.username.upper()}")
if st.sidebar.button("🚪 Logout", use_container_width=True):
    for k in ["logged_in","user_role","username"]:
        st.session_state[k] = None if k != "logged_in" else False
    st.rerun()
st.sidebar.divider()
menu = st.sidebar.radio("📋 Menu",[
    "📊 Dashboard","👤 Employee Database","📄 Contract Alerts",
    "🏖️ Leave Tracker","⚖️ Disciplinary Tracker","💰 Payroll",
    "🚨 Alerts Dashboard","📧 Notifications","📊 Reports","⚙️ Settings",
])

# ============================================================
# ========================= DASHBOARD ========================
# ============================================================
if menu == "📊 Dashboard":
    st.header("📊 HR Dashboard")
    today     = get_today()
    employees = get_employees()

    total_emp  = len(employees)
    active_emp = len(employees[employees["status"]=="active"]) if not employees.empty else 0
    left_emp   = len(employees[employees["actual_end_date"].notna()]) if not employees.empty and "actual_end_date" in employees.columns else 0

    expiring_soon_count = expired_count = 0
    if not employees.empty:
        for _,emp in employees.iterrows():
            cs = contract_status(emp.get("contract_end"),emp.get("actual_end_date"))
            if cs == "EXPIRING SOON": expiring_soon_count += 1
            elif cs == "EXPIRED":     expired_count += 1

    active_disc = expiring_disc = 0
    for d in disciplinary_col.find():
        ds = disciplinary_status(d.get("expiry_date"))
        if ds == "ACTIVE":        active_disc += 1
        elif ds == "EXPIRING SOON": expiring_disc += 1

    pending_leave = leave_col.count_documents({"approval_status":"Pending"})
    on_leave_now  = sum(1 for lv in leave_col.find({"approval_status":"Approved"})
                        if lv.get("start_date","") <= str(today) <= lv.get("end_date",""))

    st.subheader("👥 Workforce Overview")
    c1,c2,c3,c4,c5 = st.columns(5)
    c1.metric("Total Employees", total_emp)
    c2.metric("Active",          active_emp)
    c3.metric("Exited / Left",   left_emp)
    c4.metric("On Leave Today",  on_leave_now)
    c5.metric("Departments",     employees["department"].nunique() if not employees.empty else 0)
    st.divider()

    st.subheader("🚨 Pending Issues")
    p1,p2,p3,p4,p5 = st.columns(5)
    p1.metric("⏳ Expiring Contracts",  expiring_soon_count, delta="Action needed" if expiring_soon_count else None, delta_color="inverse")
    p2.metric("❌ Expired Contracts",   expired_count,       delta="Urgent"        if expired_count else None,       delta_color="inverse")
    p3.metric("⚖️ Active Disciplinary", active_disc)
    p4.metric("⚠️ Expiring Disciplinary",expiring_disc,      delta="Review soon"   if expiring_disc else None,       delta_color="inverse")
    p5.metric("📋 Pending Leave",       pending_leave,       delta="Awaiting"      if pending_leave else None,       delta_color="inverse")
    st.divider()

    col_left,col_right = st.columns(2)
    with col_left:
        st.subheader("🏢 Employees by Department")
        if not employees.empty:
            dept_summary = employees.groupby("department").agg(Total=("_id","count"),Active=("status",lambda x:(x=="active").sum())).reset_index().sort_values("Total",ascending=False)
            st.dataframe(dept_summary, use_container_width=True, hide_index=True)
        else: st.info("No data yet.")
    with col_right:
        st.subheader("📄 Contract Status Summary")
        if not employees.empty:
            statuses = [contract_status(emp.get("contract_end"),emp.get("actual_end_date")) for _,emp in employees.iterrows()]
            sc = pd.Series(statuses).value_counts().reset_index(); sc.columns = ["Status","Count"]
            st.dataframe(sc, use_container_width=True, hide_index=True)
        else: st.info("No data yet.")
    st.divider()

    st.subheader("📋 Contracts Expiring in Next 30 Days")
    expiring_list = []
    if not employees.empty:
        for _,emp in employees.iterrows():
            if contract_status(emp.get("contract_end"),emp.get("actual_end_date")) == "EXPIRING SOON":
                try:
                    end_date = datetime.strptime(str(emp["contract_end"])[:10],"%Y-%m-%d").date()
                    expiring_list.append({"Employee ID":emp.get("employee_id","—"),"Name":emp["name"],
                                          "Position":emp.get("position","—"),"Department":emp.get("department","—"),
                                          "Contract End":str(emp["contract_end"])[:10],"Days Remaining":(end_date-today).days})
                except: pass
    if expiring_list:
        st.warning(f"⚠️ {len(expiring_list)} contract(s) expiring soon!")
        st.dataframe(pd.DataFrame(expiring_list).sort_values("Days Remaining"), use_container_width=True, hide_index=True)
    else: st.success("✅ No contracts expiring in the next 30 days.")
    st.divider()

    st.subheader("🏖️ Recent Leave Requests")
    rl_docs = list(leave_col.find().sort("created_at",-1).limit(10))
    if rl_docs:
        rows = []
        for lv in rl_docs:
            emp_doc = employees_col.find_one({"_id":lv.get("employee_id")}) or {}
            rows.append({"Employee ID":emp_doc.get("employee_id","—"),"Name":emp_doc.get("name","—"),
                         "Leave Type":lv.get("leave_type",""),"Start":lv.get("start_date",""),
                         "End":lv.get("end_date",""),"Days":lv.get("days_taken",""),"Status":lv.get("approval_status","")})
        st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)
    else: st.info("No leave requests yet.")
    st.divider()

    st.subheader("⚖️ Recent Disciplinary Cases")
    rd_docs = list(disciplinary_col.find().sort("created_at",-1).limit(10))
    if rd_docs:
        rows = []
        for d in rd_docs:
            emp_doc = employees_col.find_one({"_id":d.get("employee_id")}) or {}
            rows.append({"Employee ID":emp_doc.get("employee_id","—"),"Name":emp_doc.get("name","—"),
                         "Issue Date":d.get("issue_date",""),"Issue Type":d.get("issue_type",""),
                         "Action":d.get("action_taken",""),"Expiry":d.get("expiry_date",""),
                         "Status":disciplinary_status(d.get("expiry_date"))})
        st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)
    else: st.info("No disciplinary cases yet.")


# ============================================================
# ====================== EMPLOYEE DATABASE ===================
# ============================================================
elif menu == "👤 Employee Database":
    st.header("👤 Employee Database")
    tab_list,tab_add,tab_edit,tab_delete,tab_export = st.tabs([
        "📋 All Employees","➕ Add Employee","✏️ Edit Employee","🗑️ Remove Employee","📥 Export"])

    with tab_list:
        employees = get_employees()
        if not employees.empty:
            c1,c2,c3,c4 = st.columns(4)
            c1.metric("Total",len(employees)); c2.metric("Active",len(employees[employees["status"]=="active"]))
            c3.metric("Non-Active",len(employees[employees["status"]!="active"])); c4.metric("Departments",employees["department"].nunique())
            st.divider()
        with st.expander("🔍 Filter",expanded=False):
            f1,f2,f3 = st.columns(3)
            with f1: filter_dept=st.selectbox("Department",["All"]+DEPARTMENTS)
            with f2: filter_status=st.selectbox("Status",["All"]+STATUS_OPTIONS)
            with f3: filter_contract=st.selectbox("Contract Type",["All"]+CONTRACT_TYPES)
            search_name=st.text_input("🔎 Search by name or employee ID","")
        if employees.empty:
            st.info("No employees found.")
        else:
            filtered = employees.copy()
            if filter_dept!="All":     filtered=filtered[filtered["department"]==filter_dept]
            if filter_status!="All":   filtered=filtered[filtered["status"]==filter_status]
            if filter_contract!="All": filtered=filtered[filtered["contract_type"]==filter_contract]
            if search_name:
                filtered=filtered[filtered["name"].str.contains(search_name,case=False,na=False)|
                                   filtered["employee_id"].astype(str).str.contains(search_name,case=False,na=False)]
            dcols=[c for c in ["employee_id","name","position","department","email","phone","hire_date","contract_start","contract_end","contract_type","status"] if c in filtered.columns]
            st.dataframe(filtered[dcols],use_container_width=True,hide_index=True)
            st.caption(f"Showing {len(filtered)} of {len(employees)} employees")
            if not filtered.empty:
                with st.expander("👁️ View Full Profile"):
                    sel_name = st.selectbox("Select employee",filtered["name"].tolist())
                    ed = filtered[filtered["name"]==sel_name].iloc[0]
                    c1,c2=st.columns(2)
                    with c1:
                        st.markdown("**Personal Details**")
                        for label,key in [("🪪 Employee ID","employee_id"),("👤 Name","name"),("📧 Email","email"),("📞 Phone","phone"),("🏠 Address","address")]:
                            st.write(f"{label}: {ed.get(key,'—')}")
                        st.write(f"🆘 Emergency: {ed.get('emergency_contact_name','—')} — {ed.get('emergency_contact_phone','—')}")
                    with c2:
                        st.markdown("**Employment Details**")
                        for label,key in [("💼 Position","position"),("🏢 Department","department"),("📅 Hire Date","hire_date"),("📅 Contract Start","contract_start"),("📅 Contract End","contract_end"),("📄 Contract Type","contract_type"),("🟢 Status","status")]:
                            st.write(f"{label}: {ed.get(key,'—')}")
                        if ed.get("actual_end_date"):
                            st.write(f"🚪 Actual End: {ed.get('actual_end_date','—')}")
                            st.write(f"📝 Separation: {ed.get('separation_type','—')}")

    with tab_add:
        st.subheader("➕ Add New Employee")
        with st.form("add_employee_form",clear_on_submit=True):
            st.markdown("#### 👤 Personal Information")
            a1,a2=st.columns(2)
            with a1:
                new_name=st.text_input("Full Name *"); new_email=st.text_input("Email Address")
                new_phone=st.text_input("Phone Number"); new_national_id=st.text_input("National ID")
            with a2:
                new_address=st.text_area("Address",height=68); new_emergency_name=st.text_input("Emergency Contact Name")
                new_emergency_phone=st.text_input("Emergency Contact Phone")
            st.divider(); st.markdown("#### 💼 Employment Details")
            b1,b2=st.columns(2)
            with b1:
                new_position=st.text_input("Position / Job Title *"); new_dept=st.selectbox("Department *",DEPARTMENTS)
                new_hire_date=st.date_input("Hire Date",value=get_today()); new_contract_start=st.date_input("Contract Start",value=get_today())
            with b2:
                new_contract_end=st.date_input("Contract End",value=get_today()+timedelta(days=365))
                new_contract=st.selectbox("Contract Type",CONTRACT_TYPES); new_status=st.selectbox("Status",STATUS_OPTIONS)
            st.markdown("#### 🚪 Separation (if applicable)")
            s1,s2=st.columns(2)
            with s1: new_actual_end=st.date_input("Actual End Date",value=None,min_value=date(2000,1,1))
            with s2: new_separation=st.selectbox("Separation Type",["None"]+SEPARATION_TYPES)
            new_notes=st.text_area("Notes",height=60)
            if st.form_submit_button("➕ Add Employee",use_container_width=True):
                if not new_name or not new_position:
                    st.error("Full Name and Position are required.")
                else:
                    emp_id=generate_employee_id(new_dept)
                    doc={"name":new_name,"employee_id":emp_id,"position":new_position,"department":new_dept,
                         "email":new_email,"phone":new_phone,"national_id":new_national_id,
                         "hire_date":str(new_hire_date),"contract_start":str(new_contract_start),
                         "contract_end":str(new_contract_end),"actual_end_date":str(new_actual_end) if new_actual_end else None,
                         "separation_type":new_separation if new_separation!="None" else None,
                         "contract_type":new_contract,"status":new_status,
                         "emergency_contact_name":new_emergency_name,"emergency_contact_phone":new_emergency_phone,
                         "address":new_address,"notes":new_notes,"created_at":str(get_today())}
                    try:
                        employees_col.insert_one(doc)
                        st.success(f"✅ **{new_name}** added with ID `{emp_id}`")
                        if alert_rules_col.find_one({"rule_type":"new_employee","enabled":True}):
                            plain=f"New employee:\nName:{new_name}\nPosition:{new_position}\nDept:{new_dept}\nID:{emp_id}"
                            html=build_html_email("New Employee Added",[("Name",new_name),("Position",new_position),("Dept",new_dept),("ID",emp_id)])
                            ok,msg=send_email(f"New Employee: {new_name}",plain,html_body=html)
                            if ok: st.info("📧 Notification sent.")
                    except Exception as e: st.error(f"Error: {e}")

    with tab_edit:
        employees=get_employees()
        if employees.empty: st.info("No employees to edit.")
        else:
            sel_emp=st.selectbox("Select employee",employees["name"].tolist(),key="edit_sel")
            er=employees[employees["name"]==sel_emp].iloc[0]
            with st.form("edit_employee_form"):
                e1,e2=st.columns(2)
                with e1:
                    edit_name=st.text_input("Full Name",value=er["name"]); edit_email=st.text_input("Email",value=er.get("email","") or "")
                    edit_phone=st.text_input("Phone",value=er.get("phone","") or ""); edit_position=st.text_input("Position",value=er.get("position","") or "")
                    try: cs_val=datetime.strptime(str(er.get("contract_start",""))[:10],"%Y-%m-%d").date() if er.get("contract_start") else get_today()
                    except: cs_val=get_today()
                    edit_contract_start=st.date_input("Contract Start",value=cs_val)
                    try: ce_val=datetime.strptime(str(er.get("contract_end",""))[:10],"%Y-%m-%d").date() if er.get("contract_end") else get_today()+timedelta(days=365)
                    except: ce_val=get_today()+timedelta(days=365)
                    edit_contract_end=st.date_input("Contract End",value=ce_val)
                with e2:
                    edit_dept=st.selectbox("Department",DEPARTMENTS,index=DEPARTMENTS.index(er["department"]) if er.get("department") in DEPARTMENTS else 0)
                    edit_contract=st.selectbox("Contract Type",CONTRACT_TYPES,index=CONTRACT_TYPES.index(er["contract_type"]) if er.get("contract_type") in CONTRACT_TYPES else 0)
                    edit_status=st.selectbox("Status",STATUS_OPTIONS,index=STATUS_OPTIONS.index(er["status"]) if er.get("status") in STATUS_OPTIONS else 0)
                    edit_sep=st.selectbox("Separation Type",["None"]+SEPARATION_TYPES,index=(["None"]+SEPARATION_TYPES).index(er.get("separation_type","None")) if er.get("separation_type") in SEPARATION_TYPES else 0)
                    try: ae_val=datetime.strptime(str(er.get("actual_end_date",""))[:10],"%Y-%m-%d").date() if er.get("actual_end_date") else None
                    except: ae_val=None
                    edit_actual_end=st.date_input("Actual End Date",value=ae_val,min_value=date(2000,1,1))
                edit_notes=st.text_area("Notes",value=er.get("notes","") or "")
                if st.form_submit_button("💾 Save Changes"):
                    try:
                        employees_col.update_one({"_id":ObjectId(er["_id"])},{"$set":{
                            "name":edit_name,"email":edit_email,"phone":edit_phone,"position":edit_position,
                            "department":edit_dept,"contract_start":str(edit_contract_start),"contract_end":str(edit_contract_end),
                            "actual_end_date":str(edit_actual_end) if edit_actual_end else None,
                            "separation_type":edit_sep if edit_sep!="None" else None,
                            "contract_type":edit_contract,"status":edit_status,"notes":edit_notes}})
                        st.success("✅ Employee updated."); st.rerun()
                    except Exception as e: st.error(f"Error: {e}")

    with tab_delete:
        employees=get_employees()
        if employees.empty: st.info("No employees found.")
        else:
            del_emp=st.selectbox("Select employee to remove",employees["name"].tolist())
            st.warning(f"⚠️ Remove **{del_emp}**? This cannot be undone.")
            if st.button("🗑️ Confirm Delete",type="primary"):
                employees_col.delete_one({"_id":ObjectId(employees[employees["name"]==del_emp]["_id"].iloc[0])})
                st.success(f"✅ {del_emp} removed."); st.rerun()

    with tab_export:
        employees=get_employees()
        if employees.empty: st.info("No employees to export.")
        else:
            ecols=[c for c in ["employee_id","name","position","department","email","phone","hire_date","contract_start","contract_end","actual_end_date","separation_type","contract_type","status"] if c in employees.columns]
            edf=employees[ecols]
            c1,c2=st.columns(2)
            with c1:
                xl=export_to_excel({"Employees":edf},"employees")
                st.download_button("📥 Download Excel",xl,f"employees_{datetime.now().strftime('%Y%m%d')}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
            with c2:
                st.download_button("📄 Download CSV",edf.to_csv(index=False).encode("utf-8"),f"employees_{datetime.now().strftime('%Y%m%d')}.csv",mime="text/csv",use_container_width=True)
            st.dataframe(edf,use_container_width=True,hide_index=True)


# ============================================================
# ====================== CONTRACT ALERTS =====================
# ============================================================
elif menu == "📄 Contract Alerts":
    st.header("📄 Contract Alerts")
    employees=get_employees(); today=get_today(); threshold=int(get_setting("contract_alert_days","30"))
    st.info(f"Contracts expiring within **{threshold} days**. Change in ⚙️ Settings → Payroll Defaults.")
    if employees.empty: st.warning("No employees found.")
    else:
        rows=[]
        for _,emp in employees.iterrows():
            cs=contract_status(emp.get("contract_end"),emp.get("actual_end_date"),threshold)
            try:
                end_date=datetime.strptime(str(emp.get("contract_end",""))[:10],"%Y-%m-%d").date() if emp.get("contract_end") else None
                days_rem=(end_date-today).days if end_date else None
            except: end_date=days_rem=None
            rows.append({"Employee ID":emp.get("employee_id","—"),"Name":emp["name"],"Position":emp.get("position","—"),
                         "Department":emp.get("department","—"),"Contract Start":str(emp.get("contract_start",""))[:10],
                         "Contract End":str(emp.get("contract_end",""))[:10],
                         "Actual End":str(emp.get("actual_end_date",""))[:10] if emp.get("actual_end_date") else "—",
                         "Days Remaining":days_rem,"Status":cs})
        df=pd.DataFrame(rows)
        c1,c2,c3,c4=st.columns(4)
        c1.metric("Total",len(df)); c2.metric("Active",len(df[df["Status"]=="ACTIVE"]))
        c3.metric("⏳ Expiring Soon",len(df[df["Status"]=="EXPIRING SOON"])); c4.metric("❌ Expired",len(df[df["Status"]=="EXPIRED"]))
        st.divider()
        sf=st.selectbox("Filter by Status",["All","ACTIVE","EXPIRING SOON","EXPIRED","LEFT","UNKNOWN"])
        ddf=df if sf=="All" else df[df["Status"]==sf]
        def hs(val): return {"EXPIRING SOON":"background-color:#FFF3CD;color:#856404","EXPIRED":"background-color:#F8D7DA;color:#721c24","LEFT":"background-color:#E2E3E5;color:#383d41","ACTIVE":"background-color:#D4EDDA;color:#155724"}.get(val,"")
        st.dataframe(ddf.style.applymap(hs,subset=["Status"]),use_container_width=True,hide_index=True)
        xl=export_to_excel({"Contract Alerts":ddf},"contract_alerts")
        st.download_button("📥 Export Excel",xl,f"contract_alerts_{datetime.now().strftime('%Y%m%d')}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        st.divider()
        if st.button("📧 Send Contract Expiry Alert Email"):
            expiring=df[df["Status"]=="EXPIRING SOON"]
            body=f"Contract Expiry Alert — {today}\n\n"+("\n".join([f"  - {r['Name']} ({r['Department']}) — {r['Days Remaining']} days" for _,r in expiring.iterrows()]) if not expiring.empty else "No expiring contracts.")
            html=build_html_email("Contract Expiry Alert",[(r["Name"],f"{r['Department']} — {r['Days Remaining']}d") for _,r in expiring.iterrows()] or [("Status","No expiring contracts")])
            ok,msg=send_email("HR Alert: Contract Expiry",body,html_body=html)
            st.success(msg) if ok else st.error(msg)


# ============================================================
# ====================== LEAVE TRACKER =======================
# ============================================================
elif menu == "🏖️ Leave Tracker":
    st.header("🏖️ Leave Tracker")
    employees=get_employees()
    tab_list,tab_add,tab_approve,tab_export=st.tabs(["📋 All Records","➕ Log Leave","✅ Approve / Reject","📥 Export"])

    with tab_list:
        docs=list(leave_col.find().sort("created_at",-1))
        if not docs: st.info("No leave records yet.")
        else:
            rows=[]
            for lv in docs:
                ed=employees_col.find_one({"_id":lv.get("employee_id")}) or {}
                rows.append({"Employee ID":ed.get("employee_id","—"),"Name":ed.get("name","—"),"Department":ed.get("department","—"),
                             "Leave Type":lv.get("leave_type",""),"Start":lv.get("start_date",""),"End":lv.get("end_date",""),
                             "Days":lv.get("days_taken",""),"Status":lv.get("approval_status",""),"Notes":lv.get("notes","")})
            ldf=pd.DataFrame(rows)
            c1,c2,c3,c4=st.columns(4)
            c1.metric("Total",len(ldf)); c2.metric("Pending",len(ldf[ldf["Status"]=="Pending"]))
            c3.metric("Approved",len(ldf[ldf["Status"]=="Approved"])); c4.metric("Rejected",len(ldf[ldf["Status"]=="Rejected"]))
            st.divider()
            col1,col2,col3=st.columns(3)
            with col1: fs=st.selectbox("Status",["All"]+APPROVAL_STATUS)
            with col2: ft=st.selectbox("Leave Type",["All"]+LEAVE_TYPES)
            with col3: fq=st.text_input("Search by name","")
            fl=ldf.copy()
            if fs!="All": fl=fl[fl["Status"]==fs]
            if ft!="All": fl=fl[fl["Leave Type"]==ft]
            if fq: fl=fl[fl["Name"].str.contains(fq,case=False,na=False)]
            st.dataframe(fl,use_container_width=True,hide_index=True)

    with tab_add:
        if employees.empty: st.warning("No employees found.")
        else:
            with st.form("leave_form",clear_on_submit=True):
                l1,l2=st.columns(2)
                with l1: emp_sel=st.selectbox("Employee *",employees["name"].tolist()); leave_type=st.selectbox("Leave Type *",LEAVE_TYPES); start_date=st.date_input("Start Date *",value=get_today())
                with l2: end_date=st.date_input("End Date *",value=get_today()+timedelta(days=1)); appr_status=st.selectbox("Approval Status",APPROVAL_STATUS)
                leave_notes=st.text_area("Notes",height=80)
                if st.form_submit_button("➕ Save Leave Request",use_container_width=True):
                    if end_date<start_date: st.error("End date must be after start date.")
                    else:
                        emp_doc=employees_col.find_one({"name":emp_sel})
                        days=(end_date-start_date).days
                        try:
                            leave_col.insert_one({"employee_id":emp_doc["_id"],"leave_type":leave_type,"start_date":str(start_date),"end_date":str(end_date),"days_taken":days,"approval_status":appr_status,"notes":leave_notes,"created_at":datetime.now()})
                            st.success(f"✅ Leave logged for **{emp_sel}** — {days} day(s)")
                        except Exception as e: st.error(f"Error: {e}")

    with tab_approve:
        pending_docs=list(leave_col.find({"approval_status":"Pending"}).sort("start_date",1))
        if not pending_docs: st.success("✅ No pending leave requests.")
        else:
            st.warning(f"{len(pending_docs)} pending request(s)")
            for lv in pending_docs:
                ed=employees_col.find_one({"_id":lv.get("employee_id")}) or {}
                with st.expander(f"📋 {ed.get('name','—')} — {lv.get('leave_type','')} ({lv.get('start_date','')} to {lv.get('end_date','')}, {lv.get('days_taken','')} days)"):
                    c1,c2=st.columns(2)
                    with c1:
                        if st.button("✅ Approve",key=f"approve_{lv['_id']}"):
                            leave_col.update_one({"_id":lv["_id"]},{"$set":{"approval_status":"Approved"}}); st.rerun()
                    with c2:
                        if st.button("❌ Reject",key=f"reject_{lv['_id']}"):
                            leave_col.update_one({"_id":lv["_id"]},{"$set":{"approval_status":"Rejected"}}); st.rerun()

    with tab_export:
        docs=list(leave_col.find().sort("start_date",-1))
        if not docs: st.info("No records to export.")
        else:
            rows=[]
            for lv in docs:
                ed=employees_col.find_one({"_id":lv.get("employee_id")}) or {}
                rows.append({"Employee ID":ed.get("employee_id","—"),"Name":ed.get("name","—"),"Department":ed.get("department","—"),"Leave Type":lv.get("leave_type",""),"Start":lv.get("start_date",""),"End":lv.get("end_date",""),"Days":lv.get("days_taken",""),"Status":lv.get("approval_status","")})
            edf=pd.DataFrame(rows)
            xl=export_to_excel({"Leave Records":edf},"leave")
            st.download_button("📥 Download Excel",xl,f"leave_records_{datetime.now().strftime('%Y%m%d')}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
            st.dataframe(edf,use_container_width=True,hide_index=True)


# ============================================================
# =================== DISCIPLINARY TRACKER ===================
# ============================================================
elif menu == "⚖️ Disciplinary Tracker":
    st.header("⚖️ Disciplinary Tracker")
    employees=get_employees()
    tab_list,tab_add,tab_export=st.tabs(["📋 All Cases","➕ Log New Case","📥 Export"])

    with tab_list:
        docs=list(disciplinary_col.find().sort("created_at",-1))
        if not docs: st.info("No disciplinary cases yet.")
        else:
            rows=[]
            for d in docs:
                ed=employees_col.find_one({"_id":d.get("employee_id")}) or {}
                rows.append({"Employee ID":ed.get("employee_id","—"),"Name":ed.get("name","—"),"Department":ed.get("department","—"),"Issue Date":d.get("issue_date",""),"Issue Type":d.get("issue_type",""),"Description":d.get("description",""),"Action Taken":d.get("action_taken",""),"Expiry Date":d.get("expiry_date",""),"Status (Live)":disciplinary_status(d.get("expiry_date"))})
            ddf=pd.DataFrame(rows)
            c1,c2,c3,c4=st.columns(4)
            c1.metric("Total",len(ddf)); c2.metric("⚖️ Active",len(ddf[ddf["Status (Live)"]=="ACTIVE"]))
            c3.metric("⚠️ Expiring",len(ddf[ddf["Status (Live)"]=="EXPIRING SOON"])); c4.metric("✅ Expired",len(ddf[ddf["Status (Live)"]=="EXPIRED"]))
            st.divider()
            col1,col2=st.columns(2)
            with col1: fs=st.selectbox("Status",["All","ACTIVE","EXPIRING SOON","EXPIRED","NO EXPIRY"])
            with col2: fq=st.text_input("Search by name","")
            fl=ddf.copy()
            if fs!="All": fl=fl[fl["Status (Live)"]==fs]
            if fq: fl=fl[fl["Name"].str.contains(fq,case=False,na=False)]
            def hd(val): return {"EXPIRING SOON":"background-color:#FFF3CD;color:#856404","ACTIVE":"background-color:#F8D7DA;color:#721c24","EXPIRED":"background-color:#D4EDDA;color:#155724"}.get(val,"")
            st.dataframe(fl.style.applymap(hd,subset=["Status (Live)"]),use_container_width=True,hide_index=True)

    with tab_add:
        if employees.empty: st.warning("No employees found.")
        else:
            with st.form("disc_form",clear_on_submit=True):
                d1,d2=st.columns(2)
                with d1: emp_sel=st.selectbox("Employee *",employees["name"].tolist()); issue_date=st.date_input("Issue Date *",value=get_today()); issue_type=st.selectbox("Issue Type *",ISSUE_TYPES)
                with d2: action_taken=st.text_input("Action Taken"); expiry_date=st.date_input("Warning Expiry Date",value=get_today()+timedelta(days=90))
                description=st.text_area("Description / Details *",height=100)
                if st.form_submit_button("➕ Save Case",use_container_width=True):
                    if not description: st.error("Description is required.")
                    else:
                        emp_doc=employees_col.find_one({"name":emp_sel})
                        try:
                            disciplinary_col.insert_one({"employee_id":emp_doc["_id"],"issue_date":str(issue_date),"issue_type":issue_type,"description":description,"action_taken":action_taken,"expiry_date":str(expiry_date),"status":disciplinary_status(expiry_date),"created_at":datetime.now()})
                            st.success(f"✅ Disciplinary case logged for **{emp_sel}**")
                        except Exception as e: st.error(f"Error: {e}")

    with tab_export:
        docs=list(disciplinary_col.find().sort("issue_date",-1))
        if not docs: st.info("No records to export.")
        else:
            rows=[]
            for d in docs:
                ed=employees_col.find_one({"_id":d.get("employee_id")}) or {}
                rows.append({"Employee ID":ed.get("employee_id","—"),"Name":ed.get("name","—"),"Department":ed.get("department","—"),"Issue Date":d.get("issue_date",""),"Issue Type":d.get("issue_type",""),"Description":d.get("description",""),"Action Taken":d.get("action_taken",""),"Expiry Date":d.get("expiry_date",""),"Status":disciplinary_status(d.get("expiry_date"))})
            edf=pd.DataFrame(rows)
            xl=export_to_excel({"Disciplinary Tracker":edf},"disciplinary")
            st.download_button("📥 Download Excel",xl,f"disciplinary_{datetime.now().strftime('%Y%m%d')}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
            st.dataframe(edf,use_container_width=True,hide_index=True)


# ============================================================
# ====================== PAYROLL MODULE ======================
# ============================================================
elif menu == "💰 Payroll":
    st.header("💰 Payroll Management")
    tab1,tab2,tab3,tab4=st.tabs(["Salary Settings","Run Payroll","Payroll History","Payslips"])
    employees=get_employees()

    with tab1:
        st.subheader("Set / Update Employee Salaries & Allowances")
        if employees.empty: st.warning("No employees yet.")
        else:
            emp_name=st.selectbox("Select Employee",employees["name"].tolist())
            emp_doc=employees_col.find_one({"name":emp_name})
            sett_doc=payroll_settings_col.find_one({"employee_id":emp_doc["_id"]}) if emp_doc else None
            defaults=sett_doc or {"basic_salary":0.0,"housing_allowance":0.0,"transport_allowance":0.0,"other_allowance":0.0,"tax_rate":0.0,"pension_rate":0.0,"other_deduction":0.0}
            with st.form("salary_form"):
                col1,col2=st.columns(2)
                with col1:
                    basic=st.number_input("Basic Salary (ZMW)",value=float(defaults["basic_salary"]),min_value=0.0,step=100.0)
                    housing=st.number_input("Housing Allowance",value=float(defaults["housing_allowance"]),min_value=0.0,step=50.0)
                    transport=st.number_input("Transport Allowance",value=float(defaults["transport_allowance"]),min_value=0.0,step=50.0)
                with col2:
                    other_all=st.number_input("Other Allowance",value=float(defaults["other_allowance"]),min_value=0.0,step=50.0)
                    tax_rate=st.number_input("Income Tax Rate (%)",value=float(defaults["tax_rate"])*100,min_value=0.0,max_value=50.0,step=0.5)/100
                    pension_rate=st.number_input("Pension (%)",value=float(defaults["pension_rate"])*100,min_value=0.0,max_value=20.0,step=0.5)/100
                    other_ded=st.number_input("Other Deduction (ZMW)",value=float(defaults["other_deduction"]),min_value=0.0,step=10.0)
                if st.form_submit_button("Save Salary Settings"):
                    payroll_settings_col.update_one({"employee_id":emp_doc["_id"]},{"$set":{"basic_salary":basic,"housing_allowance":housing,"transport_allowance":transport,"other_allowance":other_all,"tax_rate":tax_rate,"pension_rate":pension_rate,"other_deduction":other_ded,"currency":"ZMW","updated_at":str(get_today())}},upsert=True)
                    st.success(f"✅ Salary settings saved for {emp_name}")

    with tab2:
        st.subheader("Process Monthly Payroll")
        year,month=get_current_period()
        col1,col2=st.columns([1,3])
        with col1:
            sel_year=st.number_input("Year",min_value=2020,max_value=2035,value=year)
            sel_month=st.number_input("Month",min_value=1,max_value=12,value=month)
        existing_run=payroll_runs_col.find_one({"period_year":int(sel_year),"period_month":int(sel_month)})
        if existing_run: st.info(f"Payroll for {sel_month}/{sel_year} — Status: **{existing_run['status'].upper()}**")
        if st.button("🔄 Load / Preview Payroll"):
            if employees.empty: st.error("No employees.")
            else:
                results=[]
                for _,emp in employees.iterrows():
                    edoc=employees_col.find_one({"name":emp["name"]})
                    sdoc=payroll_settings_col.find_one({"employee_id":edoc["_id"]}) if edoc else None
                    if not sdoc: results.append({"name":emp["name"],"status":"No salary settings","net_pay":0}); continue
                    pay=calculate_pay(float(sdoc["basic_salary"]),float(sdoc["housing_allowance"]),float(sdoc["transport_allowance"]),float(sdoc["other_allowance"]),float(sdoc["tax_rate"]),float(sdoc["pension_rate"]),float(sdoc["other_deduction"]))
                    results.append({"employee_oid":str(edoc["_id"]),"name":emp["name"],"basic_salary":pay["basic_salary"],"gross_pay":pay["gross_pay"],"net_pay":pay["net_pay"],"status":"Ready"})
                preview_df=pd.DataFrame(results)
                st.session_state.payroll_preview=preview_df; st.session_state.current_run_year=int(sel_year); st.session_state.current_run_month=int(sel_month)
                st.dataframe(preview_df.style.format({"basic_salary":"ZMW {:,.2f}","gross_pay":"ZMW {:,.2f}","net_pay":"ZMW {:,.2f}"}),use_container_width=True)
                st.metric("Total Net Payroll",f"ZMW {preview_df['net_pay'].sum():,.2f}")
        if "payroll_preview" in st.session_state and st.button("💾 PROCESS Payroll"):
            try:
                sy=st.session_state.get("current_run_year",int(sel_year)); sm=st.session_state.get("current_run_month",int(sel_month))
                run_doc=payroll_runs_col.find_one({"period_year":sy,"period_month":sm})
                if run_doc: run_oid=run_doc["_id"]
                else: run_oid=payroll_runs_col.insert_one({"period_year":sy,"period_month":sm,"run_date":str(get_today()),"status":"processed"}).inserted_id
                preview=st.session_state.payroll_preview; total_net=0
                for _,row in preview.iterrows():
                    if row["status"]!="Ready": continue
                    emp_oid=ObjectId(row["employee_oid"])
                    sdoc=payroll_settings_col.find_one({"employee_id":emp_oid})
                    if not sdoc: continue
                    pay=calculate_pay(float(sdoc["basic_salary"]),float(sdoc["housing_allowance"]),float(sdoc["transport_allowance"]),float(sdoc["other_allowance"]),float(sdoc["tax_rate"]),float(sdoc["pension_rate"]),float(sdoc["other_deduction"]))
                    total_net+=pay["net_pay"]
                    payroll_payments_col.insert_one({"run_id":run_oid,"employee_id":emp_oid,"basic_salary":pay["basic_salary"],"allowances":pay["allowances"],"gross_pay":pay["gross_pay"],"tax":pay["tax"],"pension":pay["pension"],"other_deduction":pay["other_deduction"],"net_pay":pay["net_pay"],"payment_date":str(get_today()),"status":"pending"})
                st.success(f"✅ Payroll processed for {sm}/{sy}!"); st.balloons(); del st.session_state.payroll_preview
            except Exception as e: st.error(f"Error: {e}")

    with tab3:
        runs=list(payroll_runs_col.find().sort([("period_year",-1),("period_month",-1)]))
        if not runs: st.info("No payroll runs found.")
        else:
            runs_df=pd.DataFrame([{"Period":f"{r['period_year']}-{r['period_month']:02d}","Run Date":r.get("run_date",""),"Status":r.get("status","")} for r in runs])
            st.dataframe(runs_df,use_container_width=True,hide_index=True)
            vp=st.selectbox("View details",runs_df["Period"].tolist())
            sel_run=next((r for r in runs if f"{r['period_year']}-{r['period_month']:02d}"==vp),None)
            if sel_run:
                payments=list(payroll_payments_col.find({"run_id":sel_run["_id"]}))
                if payments:
                    pay_rows=[{"Name":(employees_col.find_one({"_id":p["employee_id"]}) or {}).get("name","—"),"Basic Salary":p.get("basic_salary",0),"Gross Pay":p.get("gross_pay",0),"Net Pay":p.get("net_pay",0),"Status":p.get("status","")} for p in payments]
                    st.dataframe(pd.DataFrame(pay_rows).style.format({"Basic Salary":"ZMW {:,.2f}","Gross Pay":"ZMW {:,.2f}","Net Pay":"ZMW {:,.2f}"}),use_container_width=True,hide_index=True)

    with tab4:
        runs=list(payroll_runs_col.find().sort([("period_year",-1),("period_month",-1)]))
        if not runs: st.info("No payroll runs available.")
        else:
            labels=[f"{r['period_year']}-{r['period_month']:02d}" for r in runs]
            sel_period=st.selectbox("Select Payroll Run",labels)
            sel_run=next((r for r in runs if f"{r['period_year']}-{r['period_month']:02d}"==sel_period),None)
            if sel_run:
                payments=list(payroll_payments_col.find({"run_id":sel_run["_id"]}))
                slip_rows=[{**p,"name":(employees_col.find_one({"_id":p["employee_id"]}) or {}).get("name","—"),"position":(employees_col.find_one({"_id":p["employee_id"]}) or {}).get("position","—"),"department":(employees_col.find_one({"_id":p["employee_id"]}) or {}).get("department","—"),"period":sel_period} for p in payments]
                if not slip_rows: st.warning("No payslips for this run.")
                else:
                    emp_choice=st.selectbox("Employee Payslip",[s["name"] for s in slip_rows])
                    slip=next((s for s in slip_rows if s["name"]==emp_choice),None)
                    if slip:
                        st.markdown(f"""**{get_setting('company_name','Kaunda HRMS')} — Payslip {slip['period']}**
**{slip['name']}** — {slip['position']} | {slip['department']}

**Earnings**
Basic Salary: ZMW {slip['basic_salary']:,.2f}
Allowances:   ZMW {slip['allowances']:,.2f}
**Gross Pay**: ZMW **{slip['gross_pay']:,.2f}**

**Deductions**
PAYE/Tax: ZMW {slip['tax']:,.2f}  Pension: ZMW {slip['pension']:,.2f}  Other: ZMW {slip['other_deduction']:,.2f}

**Net Pay**: ZMW **{slip['net_pay']:,.2f}**""")
                        all_df=pd.DataFrame([{k:v for k,v in s.items() if k not in ["_id","run_id","employee_id"]} for s in slip_rows])
                        xl=export_to_excel({"Payslips":all_df},f"payslips_{sel_period}")
                        st.download_button("📥 Export All Payslips",xl,f"payslips_{sel_period}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ============================================================
# ====================== ALERTS DASHBOARD ====================
# ============================================================
elif menu == "🚨 Alerts Dashboard":
    st.header("🚨 Alerts Dashboard")
    employees=get_employees(); today=get_today()
    col1,col2,col3=st.columns(3)
    expiring=[]
    with col1:
        st.subheader("📄 Contract Expiry")
        threshold=int(get_setting("contract_alert_days","30"))
        if not employees.empty:
            for _,emp in employees.iterrows():
                if contract_status(emp.get("contract_end"),emp.get("actual_end_date"),threshold)=="EXPIRING SOON":
                    try:
                        end=datetime.strptime(str(emp["contract_end"])[:10],"%Y-%m-%d").date()
                        expiring.append({"Name":emp["name"],"Dept":emp["department"],"Days Left":(end-today).days})
                    except: pass
        if expiring: st.warning(f"{len(expiring)} contract(s) expiring soon!"); st.dataframe(pd.DataFrame(expiring),use_container_width=True,hide_index=True)
        else: st.success("✅ No contracts expiring soon.")
    active_cases=[]
    with col2:
        st.subheader("⚖️ Active Disciplinary Cases")
        for d in disciplinary_col.find():
            ds=disciplinary_status(d.get("expiry_date"))
            if ds in ["ACTIVE","EXPIRING SOON"]:
                ed=employees_col.find_one({"_id":d.get("employee_id")}) or {}
                active_cases.append({"Name":ed.get("name","—"),"Issue":d.get("issue_type",""),"Status":ds})
        if active_cases: st.warning(f"{len(active_cases)} active case(s)"); st.dataframe(pd.DataFrame(active_cases),use_container_width=True,hide_index=True)
        else: st.success("✅ No active disciplinary cases.")
    pending_rows=[]
    with col3:
        st.subheader("🏖️ Pending Leave Requests")
        for lv in leave_col.find({"approval_status":"Pending"}):
            ed=employees_col.find_one({"_id":lv.get("employee_id")}) or {}
            pending_rows.append({"Name":ed.get("name","—"),"Leave Type":lv.get("leave_type",""),"Start":lv.get("start_date",""),"End":lv.get("end_date","")})
        if pending_rows: st.warning(f"{len(pending_rows)} pending request(s)"); st.dataframe(pd.DataFrame(pending_rows),use_container_width=True,hide_index=True)
        else: st.success("✅ No pending leave requests.")
    st.divider()
    st.subheader("📧 Send Alert Email Now")
    alert_type=st.selectbox("Alert Type",["Contract Expiry Summary","Disciplinary Summary","Leave Summary","Custom Alert"])
    if alert_type=="Custom Alert":
        cs=st.text_input("Subject"); cb=st.text_area("Body",height=120)
        if st.button("📤 Send"):
            ok,msg=send_email(cs,cb); st.success(msg) if ok else st.error(msg)
    else:
        if st.button(f"📤 Send {alert_type} Email"):
            if alert_type=="Contract Expiry Summary":
                body=f"Contract Expiry — {today}\n\n"+("\n".join([f"  - {e['Name']} ({e['Dept']}) — {e['Days Left']} days" for e in expiring]) or "None.")
                html=build_html_email("Contract Expiry",[(e["Name"],f"{e['Dept']} — {e['Days Left']}d") for e in expiring] or [("Status","None")])
            elif alert_type=="Disciplinary Summary":
                body=f"Disciplinary Summary — {today}\n\n"+("\n".join([f"  - {c['Name']} ({c['Issue']})" for c in active_cases]) or "None.")
                html=build_html_email("Disciplinary Summary",[(c["Name"],c["Issue"]) for c in active_cases] or [("Status","None")])
            else:
                body=f"Pending Leave — {today}\n\n"+("\n".join([f"  - {r['Name']} ({r['Leave Type']}): {r['Start']} to {r['End']}" for r in pending_rows]) or "None.")
                html=build_html_email("Leave Summary",[(r["Name"],r["Leave Type"]) for r in pending_rows] or [("Status","None")])
            ok,msg=send_email(f"HR Alert: {alert_type}",body,html_body=html)
            st.success(msg) if ok else st.error(msg)


# ============================================================
# ====================== NOTIFICATIONS =======================
# ============================================================
elif menu == "📧 Notifications":
    st.header("📧 Notifications Centre")
    tab_send,tab_rules,tab_log=st.tabs(["✉️ Send Notification","⚙️ Alert Rules","📋 Notification Log"])
    with tab_send:
        employees=get_employees()
        notif_type=st.radio("Recipient",["Admin Email","Specific Employee","All Active Employees"],horizontal=True)
        with st.form("notification_form"):
            subject=st.text_input("Subject *"); message=st.text_area("Message *",height=150)
            if notif_type=="Specific Employee": emp_sel=st.selectbox("Select Employee",employees["name"].tolist() if not employees.empty else ["No employees"])
            include_html=st.checkbox("Send as formatted HTML email",value=True)
            if st.form_submit_button("📤 Send Notification",use_container_width=True):
                if not subject or not message: st.error("Subject and message are required.")
                else:
                    if notif_type=="Admin Email": recipients=[get_setting("admin_email","")]
                    elif notif_type=="Specific Employee":
                        ed=employees_col.find_one({"name":emp_sel}) or {}; recipients=[ed.get("email","")] if ed.get("email") else []
                    else: recipients=[e for e in employees[employees["status"]=="active"]["email"].dropna().tolist() if e] if not employees.empty else []
                    sent=failed=0
                    for r in recipients:
                        html_body=build_html_email(subject,[("Message",message)]) if include_html else None
                        ok,_=send_email(subject,message,recipient=r,html_body=html_body)
                        if ok: sent+=1
                        else: failed+=1
                    if sent>0: st.success(f"✅ Sent to {sent} recipient(s).")
                    if failed>0: st.error(f"❌ Failed for {failed} recipient(s).")
    with tab_rules:
        for rule in alert_rules_col.find():
            col1,col2,col3=st.columns([3,1,1])
            with col1: st.write(f"**{rule['rule_name']}**  \n`{rule['rule_type']}`")
            with col2:
                if rule.get("threshold_days",0)>0: st.write(f"⏰ {rule['threshold_days']} days")
            with col3:
                if st.button("✅ Enabled" if rule.get("enabled") else "❌ Disabled",key=f"toggle_{rule['_id']}"):
                    alert_rules_col.update_one({"_id":rule["_id"]},{"$set":{"enabled":not rule.get("enabled",True)}}); st.rerun()
    with tab_log:
        log_docs=list(notif_col.find().sort("sent_at",-1).limit(100))
        if not log_docs: st.info("No notifications sent yet.")
        else:
            st.dataframe(pd.DataFrame([{"Sent At":str(l.get("sent_at","")),"Subject":l.get("subject",""),"Recipient":l.get("recipient",""),"Status":l.get("status","")} for l in log_docs]),use_container_width=True,hide_index=True)
            if st.button("🗑️ Clear Log"):
                notif_col.delete_many({}); st.success("Log cleared."); st.rerun()


# ============================================================
# ====================== REPORTS =============================
# ============================================================
elif menu == "📊 Reports":
    st.header("📊 Reports")
    employees=get_employees()
    if not employees.empty:
        c1,c2,c3,c4=st.columns(4)
        c1.metric("Total Employees",len(employees)); c2.metric("Active",len(employees[employees["status"]=="active"]))
        c3.metric("Departments",employees["department"].nunique()); c4.metric("Contract Types",employees["contract_type"].nunique())
        st.subheader("Department Summary")
        st.dataframe(employees.groupby("department").agg(Total=("_id","count"),Active=("status",lambda x:(x=="active").sum())).reset_index(),use_container_width=True)
        st.subheader("Contract Type Breakdown")
        ct=employees["contract_type"].value_counts().reset_index(); ct.columns=["Contract Type","Count"]
        st.dataframe(ct,use_container_width=True)
        st.subheader("Leave Summary by Type")
        ldocs=list(leave_col.find())
        if ldocs:
            ldf=pd.DataFrame(ldocs)
            st.dataframe(ldf.groupby("leave_type").agg(Requests=("_id","count"),Total_Days=("days_taken","sum"),Approved=("approval_status",lambda x:(x=="Approved").sum())).reset_index(),use_container_width=True)
        st.subheader("Disciplinary Summary by Type")
        ddocs=list(disciplinary_col.find())
        if ddocs:
            ddf=pd.DataFrame(ddocs)
            st.dataframe(ddf.groupby("issue_type").agg(Cases=("_id","count")).reset_index(),use_container_width=True)
        st.subheader("Payroll Summary")
        runs=list(payroll_runs_col.find().sort([("period_year",-1),("period_month",-1)]))
        if runs:
            pay_summary=[{"Period":f"{r['period_year']}-{r['period_month']:02d}","Employees Paid":payroll_payments_col.count_documents({"run_id":r["_id"]}),"Total Net Pay":sum(p.get("net_pay",0) for p in payroll_payments_col.find({"run_id":r["_id"]})),"Status":r.get("status","")} for r in runs]
            st.dataframe(pd.DataFrame(pay_summary).style.format({"Total Net Pay":"ZMW {:,.2f}"}),use_container_width=True)
    else: st.info("No employee data available yet.")


# ============================================================
# ======================== SETTINGS ==========================
# ============================================================
elif menu == "⚙️ Settings":
    st.header("⚙️ System Settings")
    tab_admin,tab_company,tab_email,tab_payroll_cfg,tab_security=st.tabs(["🔐 Admin Credentials","🏢 Company Info","📧 Email Configuration","💰 Payroll Defaults","🛡️ Security"])

    with tab_admin:
        st.subheader("🔐 Change Admin Credentials")
        with st.form("admin_cred_form"):
            current_user=get_setting("admin_username","admin"); current_pass=get_setting("admin_password","admin123")
            st.info(f"Current Username: **{current_user}**")
            new_username=st.text_input("New Username"); new_password=st.text_input("New Password",type="password")
            confirm_pass=st.text_input("Confirm Password",type="password"); verify_current=st.text_input("Current Password (to confirm)",type="password")
            if st.form_submit_button("💾 Update Credentials",use_container_width=True):
                if verify_current!=current_pass: st.error("❌ Current password incorrect.")
                elif new_password and new_password!=confirm_pass: st.error("❌ Passwords do not match.")
                elif new_password and len(new_password)<6: st.error("Password must be at least 6 characters.")
                else:
                    if new_username: save_setting("admin_username",new_username)
                    if new_password: save_setting("admin_password",new_password)
                    st.success("✅ Credentials updated. Please log in again.")
                    st.session_state.logged_in=False; st.rerun()

    with tab_company:
        st.subheader("🏢 Company Information")
        with st.form("company_form"):
            cn=st.text_input("Company Name",value=get_setting("company_name","Kaunda HRMS"))
            ca=st.text_area("Company Address",value=get_setting("company_address","Lusaka, Zambia"),height=80)
            cp=st.text_input("Company Phone",value=get_setting("company_phone",""))
            ce=st.text_input("Company Email",value=get_setting("company_email",""))
            cw=st.text_input("Website",value=get_setting("company_website",""))
            ct=st.text_input("TPIN",value=get_setting("company_tpin",""))
            cn2=st.text_input("NAPSA No.",value=get_setting("company_napsa",""))
            cc=st.selectbox("Default Currency",["ZMW","USD","ZAR","EUR","GBP"],index=["ZMW","USD","ZAR","EUR","GBP"].index(get_setting("default_currency","ZMW")))
            if st.form_submit_button("💾 Save",use_container_width=True):
                for k,v in [("company_name",cn),("company_address",ca),("company_phone",cp),("company_email",ce),("company_website",cw),("company_tpin",ct),("company_napsa",cn2),("default_currency",cc)]: save_setting(k,v)
                st.success("✅ Saved.")

    with tab_email:
        st.subheader("📧 Email / SMTP Configuration")
        with st.form("email_config_form"):
            col1,col2=st.columns(2)
            with col1:
                smtp_host=st.text_input("SMTP Host",value=get_setting("smtp_host","smtp.gmail.com"))
                smtp_port=st.selectbox("SMTP Port",["587","465","25"],index=["587","465","25"].index(get_setting("smtp_port","587")))
                smtp_user=st.text_input("Email",value=get_setting("smtp_user",""))
            with col2:
                smtp_pass=st.text_input("App Password",type="password",value=get_setting("smtp_password",""))
                admin_email=st.text_input("Admin Recipient Email",value=get_setting("admin_email",""))
            st.markdown("**Gmail users:** Use an App Password. Enable 2FA first.")
            if st.form_submit_button("💾 Save Email Settings",use_container_width=True):
                for k,v in [("smtp_host",smtp_host),("smtp_port",smtp_port),("smtp_user",smtp_user),("smtp_password",smtp_pass),("admin_email",admin_email)]: save_setting(k,v)
                st.success("✅ Email settings saved.")
        st.divider()
        test_rec=st.text_input("Send test to",value=get_setting("admin_email",""))
        if st.button("📤 Send Test Email"):
            ok,msg=send_email("HRMS Test Email",f"Test email. Sent: {datetime.now().strftime('%d %b %Y %H:%M')}",recipient=test_rec)
            st.success(msg) if ok else st.error(msg)

    with tab_payroll_cfg:
        st.subheader("💰 Payroll Defaults")
        with st.form("payroll_defaults_form"):
            col1,col2=st.columns(2)
            with col1:
                dt=st.number_input("Default PAYE Tax Rate (%)",value=float(get_setting("default_tax_rate","30")),min_value=0.0,max_value=60.0,step=0.5)
                dp=st.number_input("Default Pension Rate (%)",value=float(get_setting("default_pension_rate","5")),min_value=0.0,max_value=20.0,step=0.5)
            with col2:
                ne=st.number_input("NAPSA Employee (%)",value=float(get_setting("napsa_employee","5")),min_value=0.0,max_value=10.0,step=0.5)
                nr=st.number_input("NAPSA Employer (%)",value=float(get_setting("napsa_employer","5")),min_value=0.0,max_value=10.0,step=0.5)
            cad=st.number_input("Contract Expiry Alert Threshold (days)",value=int(get_setting("contract_alert_days","30")),min_value=1,max_value=180)
            if st.form_submit_button("💾 Save",use_container_width=True):
                for k,v in [("default_tax_rate",str(dt)),("default_pension_rate",str(dp)),("napsa_employee",str(ne)),("napsa_employer",str(nr)),("contract_alert_days",str(cad))]: save_setting(k,v)
                st.success("✅ Saved.")

    with tab_security:
        st.subheader("🛡️ Security Settings")
        with st.form("security_form"):
            al=st.selectbox("Auto Logout",["Never","15 minutes","30 minutes","1 hour","2 hours"],index=["Never","15 minutes","30 minutes","1 hour","2 hours"].index(get_setting("auto_logout","Never")))
            rs=st.checkbox("Require Strong Passwords",value=get_setting("strong_password","0")=="1")
            if st.form_submit_button("💾 Save"): save_setting("auto_logout",al); save_setting("strong_password","1" if rs else "0"); st.success("✅ Saved.")
        st.divider()
        st.subheader("System Info")
        c1,c2=st.columns(2)
        with c1:
            st.metric("Total Employees",employees_col.count_documents({}))
            st.metric("Payroll Runs",payroll_runs_col.count_documents({}))
            st.metric("Notifications Sent",notif_col.count_documents({}))
        with c2:
            st.metric("Leave Records",leave_col.count_documents({}))
            st.metric("Disciplinary Cases",disciplinary_col.count_documents({}))
            st.metric("Company",get_setting("company_name","Kaunda HRMS"))

# ====================== FOOTER ======================
st.sidebar.divider()
st.sidebar.caption(f"🏢 {get_setting('company_name','Kaunda HRMS')}")
st.sidebar.caption("☁️ Data stored in MongoDB Atlas")
st.sidebar.caption(f"🔐 {st.session_state.username.upper()}")